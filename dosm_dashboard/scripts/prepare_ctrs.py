"""
Extract the precomputed CTRS/TSRI result (coastal tourism risk index) from
the TSRI_CTRS_model.xlsx workbook into a clean CSV the app reads directly.
The app does NOT recompute the index — it trusts CTRS_Summary as the
source of truth, so this script is a straight read + state-name-canon step.

Run: python dosm_dashboard/scripts/prepare_ctrs.py
"""
import os

import openpyxl
import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
NB = os.path.join(ROOT, "notebooks")
OUT = os.path.join(ROOT, "data", "processed")

CANON = {
    "Johor": "Johor", "Kedah": "Kedah", "Kelantan": "Kelantan", "Melaka": "Malacca",
    "Negeri Sembilan": "Negeri Sembilan", "Pulau Pinang": "Penang", "Pahang": "Pahang",
    "Perak": "Perak", "Perlis": "Perlis", "Sabah": "Sabah", "Sarawak": "Sarawak",
    "Selangor": "Selangor", "Terengganu": "Terengganu", "W.P. Labuan": "Labuan",
    "W.P. Kuala Lumpur": "Kuala Lumpur", "W.P. Putrajaya": "Putrajaya",
}


def main():
    src = os.path.join(NB, "TSRI_CTRS_model.xlsx")
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)

    ws = wb["CTRS_Summary"]
    rows = list(ws.iter_rows(values_only=True))
    header = rows[4]
    data_rows = [r for r in rows[5:] if isinstance(r[0], (int, float))]
    df = pd.DataFrame(data_rows, columns=header)
    df = df.rename(columns={
        "Rank": "rank", "State": "state", "Pantai share": "pantai_share",
        "X": "x_exposure", "Y": "y_stress", "Z": "z_protection",
        "TSRI": "tsri", "TSRI (0–1)": "tsri_0_1", "LCC note": "lcc_note",
    })
    df["state_canon"] = df["state"].map(lambda s: CANON.get(s, s))
    df.to_csv(os.path.join(OUT, "ctrs_summary.csv"), index=False)
    print(f"[ctrs] {len(df)} states -> {os.path.join(OUT, 'ctrs_summary.csv')}")

    # Key control-panel settings, for an in-app "model settings used" note.
    ctrl = wb["CTRS_Control"]
    ctrl_rows = {r[1]: r[2] for r in ctrl.iter_rows(values_only=True) if r and r[1]}
    settings = {
        "x_method": ctrl_rows.get("X_METHOD_IN_USE", ctrl_rows.get("X_METHOD")),
        "mmwqi_year": ctrl_rows.get("MMWQI_YEAR"),
        "tourist_year": ctrl_rows.get("TOURIST_YEAR"),
        "lcc_missing": ctrl_rows.get("LCC_MISSING"),
        "z_method": ctrl_rows.get("Z_METHOD"),
        "z_denom": ctrl_rows.get("Z_DENOM"),
        "w_x": ctrl_rows.get("w_X"), "w_y": ctrl_rows.get("w_Y"), "w_z": ctrl_rows.get("w_Z"),
        "w_mmwqi": ctrl_rows.get("w_MMWQI"), "w_lcc": ctrl_rows.get("w_LCC"),
    }
    pd.Series(settings).to_csv(os.path.join(OUT, "ctrs_settings.csv"), header=["value"])
    print(f"[ctrs] settings -> {os.path.join(OUT, 'ctrs_settings.csv')}")


if __name__ == "__main__":
    main()
